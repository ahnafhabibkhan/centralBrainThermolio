"""Bounded extraction subprocess. Never executes document macros or formulas."""

import csv
import json
import resource
import sys
import zipfile
from pathlib import Path


def extract(path, extension):
    sections = []
    total = 0
    truncated = False

    def add(location, text):
        nonlocal total, truncated
        text = text.replace("\x00", "").strip()
        for start in range(0, len(text), 4000):
            if total >= 1000000 or len(sections) >= 500:
                truncated = True
                return False
            content = text[start : start + min(4000, 1000000 - total)]
            sections.append({"location": location, "content": content})
            total += len(content)
        return True

    if extension in {".docx", ".xlsx"}:
        with zipfile.ZipFile(path) as archive:
            if (
                sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024
                or len(archive.infolist()) > 10000
            ):
                raise ValueError("Expanded document exceeds processing limits.")
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs cannot be indexed.")
        for index, page in enumerate(reader.pages):
            if index >= 500:
                truncated = True
                break
            if not add(f"Page {index + 1}", page.extract_text() or ""):
                break
    elif extension == ".docx":
        from defusedxml.ElementTree import fromstring

        with zipfile.ZipFile(path) as archive:
            root = fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for index, p in enumerate(root.findall(".//w:p", ns)):
            if not add(
                f"Paragraph {index + 1}", "".join(t.text or "" for t in p.findall(".//w:t", ns))
            ):
                break
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        for sheet in workbook:
            if not add(f"Sheet {sheet.title}", f"Worksheet: {sheet.title}"):
                break
            for index, row in enumerate(sheet.iter_rows(max_col=100, values_only=True)):
                if index >= 10000:
                    truncated = True
                    break
                text = " | ".join("" if cell is None else str(cell)[:2000] for cell in row)
                if not add(f"Sheet {sheet.title}, row {index + 1}", text):
                    break
            if total >= 1000000 or len(sections) >= 500:
                break
        workbook.close()
    elif extension == ".csv":
        csv.field_size_limit(100000)
        with open(path, encoding="utf-8-sig") as stream:
            for index, row in enumerate(csv.reader(stream)):
                if not add(f"Row {index + 1}", " | ".join(row)):
                    break
    else:
        with open(path, encoding="utf-8-sig") as stream:
            text = stream.read(1000001)
        if len(text) > 1000000:
            truncated = True
        add("Text", text)
    return {
        "sections": sections,
        "state": "partial" if truncated else ("ready" if sections else "unsearchable"),
        "error": "Extraction reached the pilot limit. Download the original for remaining content."
        if truncated
        else (
            None
            if sections
            else "No readable text was found. Scanned PDFs need OCR, which is not enabled."
        ),
    }


if __name__ == "__main__":
    if sys.platform == "linux":
        resource.setrlimit(resource.RLIMIT_AS, (384 * 1024 * 1024, 384 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (45, 45))
    resource.setrlimit(resource.RLIMIT_FSIZE, (8 * 1024 * 1024, 8 * 1024 * 1024))
    try:
        result = extract(sys.argv[1], sys.argv[2])
    except Exception:
        result = {
            "sections": [],
            "state": "failed",
            "error": "The file could not be processed. It may be encrypted, malformed, non-UTF-8 text, or exceed processing limits. The original remains downloadable.",
        }
    Path(sys.argv[3]).write_text(json.dumps(result))
