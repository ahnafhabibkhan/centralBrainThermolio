import io
import re
import zipfile
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from test_pilot import pilot, memory
from central_brain.auth import AuthContext
from central_brain.extract import extract
from central_brain.library import Library
from central_brain.library_worker import process_one


@pytest.fixture
def library(pilot, tmp_path):
    pilot.settings.library_local_path = str(tmp_path)
    return Library(pilot.repo, pilot.settings)


def test_upload_version_search_and_original(pilot, library):
    auth = pilot.auth
    folder = library.folder(auth, "180 Bloor")
    original = b"Heating decision: install efficient boiler controls."
    node = library.upload(auth, "Decision.txt", original, folder)
    assert library.read(auth, node)["state"] == "queued"
    assert process_one(library, auth)
    hit = library.search(auth, "boiler", folder)["results"][0]
    assert hit["id"] == node and hit["path"] == "/180 Bloor/Decision.txt"
    name, stream = library.download(auth, node)
    with stream:
        assert stream.read() == original
    library.upload(auth, "Decision.txt", b"Cooling upgrade approved.", node_id=node)
    assert not library.search(auth, "boiler")["results"]
    assert process_one(library, auth)
    assert library.search(auth, "cooling")["results"][0]["version"] == 2
    assert library.read(auth, node, version=1)["sections"][0]["content"] == original.decode()
    assert library.usage(auth)["used"] == len(original) + len(b"Cooling upgrade approved.")
    library.move(auth, node, "Decision-renamed.txt", None)
    assert library.info(auth, node)["path"] == "/Decision-renamed.txt"
    assert not library.search(auth, "cooling", folder)["results"]


def test_private_proposals_and_memory_projection(pilot, library):
    auth = pilot.auth
    other = AuthContext(auth.principal.model_copy(update={"actor_id": uuid4()}))
    node = library.upload(auth, "Private.txt", b"Secret building notes.", visibility="private")
    assert process_one(library, auth)
    assert library.listing(other) == []
    assert library.search(other, "building")["results"] == []
    with pytest.raises(HTTPException):
        library.read(other, node)
    proposed = library.upload(auth, "Proposal.md", b"Proposed preference.", proposed=True)
    assert process_one(library, auth)
    with pytest.raises(HTTPException):
        library.read(auth, proposed)
    assert library.search(auth, "preference")["results"] == []
    library.review(auth, proposed, True)
    assert library.search(auth, "preference")["results"]
    item = pilot.repo.create(auth, memory())
    with pytest.raises(HTTPException):
        library.read(auth, item.memory_id)
    pilot.repo.approve(auth, item.memory_id)
    assert library.read(auth, item.memory_id)["sections"]
    library.move(auth, item.memory_id, "Preference.md", None)
    pilot.repo.transition(auth, item.memory_id, "delete")
    with pytest.raises(HTTPException):
        library.read(auth, item.memory_id)


def test_limits_cycles_and_quota_race(pilot, library):
    auth = pilot.auth
    a = library.folder(auth, "A")
    b = library.folder(auth, "B", a)
    with pytest.raises(HTTPException):
        library.move(auth, a, "A", b)
    with pytest.raises(HTTPException):
        library.folder(auth, "A")
    with pytest.raises(HTTPException):
        library.upload(auth, "../escape.txt", b"x")
    with pytest.raises(HTTPException):
        library.upload(auth, "macro.xlsm", b"x")
    pilot.settings.library_file_bytes = 10
    with pytest.raises(HTTPException):
        library.upload(auth, "big.txt", b"x" * 11)
    pilot.settings.library_quota_bytes = 10

    def put(index):
        try:
            library.upload(auth, f"{index}.txt", b"x" * 8)
            return True
        except HTTPException:
            return False

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sum(pool.map(put, range(2))) == 1


def test_extract_formats_and_bounds(tmp_path):
    text = tmp_path / "text"
    text.write_text("x" * 1001000)
    assert extract(text, ".txt")["state"] == "partial"
    text.write_text("name,value\nboiler,42\n")
    assert extract(text, ".csv")["sections"][1]["location"] == "Row 2"
    doc = tmp_path / "doc.docx"
    with zipfile.ZipFile(doc, "w") as z:
        z.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Heating decision</w:t></w:r></w:p></w:body></w:document>',
        )
    assert extract(doc, ".docx")["sections"][0]["content"] == "Heating decision"
    sheet = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Costs"
    ws.append(["Equipment", "CAD"])
    ws.append(["Boiler", 42])
    wb.save(sheet)
    assert any(s["location"] == "Sheet Costs, row 2" for s in extract(sheet, ".xlsx")["sections"])
    from pypdf import PdfWriter

    pdf = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf)
    assert extract(pdf, ".pdf")["state"] == "unsearchable"


def test_web_upload_csrf_preview_and_download(pilot, library):
    client = pilot.client
    login = client.get("/login").text
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', login)[1]
    client.post("/login", data={"token": "owner", "csrf_token": csrf})
    html = client.get("/library").text
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', html)[1]
    assert (
        client.post(
            "/library/upload",
            files={"file": ("Sample.txt", b"Boiler test")},
            data={"csrf_token": "wrong"},
        ).status_code
        == 403
    )
    response = client.post(
        "/library/upload",
        files={"file": ("Sample.txt", b"Boiler test")},
        data={"csrf_token": csrf},
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text
    url = response.headers["location"]
    assert client.get(url).status_code == 200
    assert client.get(url + "/download").content == b"Boiler test"


def test_library_navigation_and_markdown_memory_replacements(pilot, library):
    client = pilot.client
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', client.get('/login').text)[1]
    client.post('/login', data={'token': 'owner', 'csrf_token': csrf})
    roots, folders = library.workspace_folders(pilot.auth)
    assert {'Memories', 'Skills'} == set(roots)
    assert library.workspace_folders(pilot.auth)[0] == roots
    page = client.get('/library').text
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', page)[1]
    assert 'Folder hierarchy' in page and 'account-actions' in page
    assert 'folder-sidebar' not in page and '>All files</a>' not in page
    assert 'No approvals are waiting.' in page
    assert 'data-tree-id=' in page and 'Memories' in page and 'Skills' in page
    assert client.get('/skills', follow_redirects=False).headers['location'] == f"/library?folder={roots['Skills']}"
    data = {'csrf_token': csrf, 'memory_type': 'fact', 'source_reference': 'Test file'}
    assert client.post('/new', data=data, files={'file': ('bad.txt', b'Fact')}).status_code == 422
    created = client.post('/new', data=data, files={'file': ('note.md', b'# Test\nDisplay costs in CAD.')}, follow_redirects=False)
    assert created.status_code == 303
    url = created.headers['location']
    download = client.get(url + '/export')
    assert '.md' in download.headers['content-disposition']
    assert download.content == b'# Test\nDisplay costs in CAD.'
    client.post(url + '/approve', data={'csrf_token': csrf})
    updated = client.post(url + '/edit', data=data, files={'file': ('note.md', b'# Updated\nDisplay costs in CAD with tax.')}, follow_redirects=False)
    assert updated.status_code == 303
    assert client.get(url + '/export').content == download.content
    original_id = url.rsplit('/', 1)[1]
    from uuid import UUID
    assert pilot.repo.get(pilot.auth, UUID(original_id)).status == 'active'
    client.post(updated.headers['location'] + '/approve', data={'csrf_token': csrf})
    assert pilot.repo.get(pilot.auth, UUID(original_id)).status == 'superseded'
    uploaded = client.post('/library/upload', data={'csrf_token': csrf, 'parent_id': str(roots['Memories'])},
                           files={'file': ('folder-note.md', b'Memory from the folder.')}, follow_redirects=False)
    assert uploaded.status_code == 303 and uploaded.headers['location'].startswith('/review/')
    assert client.post('/library/upload', data={'csrf_token': csrf, 'parent_id': str(roots['Memories'])},
                       files={'file': ('bad.pdf', b'Not Markdown')}).status_code == 422
    skill = library.upload(pilot.auth, 'Procedure.md', b'Check the boiler sequence.', roots['Skills'])
    assert 'Procedure.md' in client.get('/skills').text
    assert process_one(library, pilot.auth)
    assert library.search(pilot.auth, 'boiler', roots['Skills'])['results'][0]['id'] == skill
    proposed = library.upload(pilot.auth, 'Humanizer.md', b'Improve the writing.', roots['Skills'], proposed=True)
    skill_page = client.get('/skills').text
    assert skill_page.index('id="approvals"') < skill_page.index('id="folder-content"')
    assert 'Humanizer.md' in skill_page


def test_context_changes_and_skill_catalog_respect_access(pilot, library):
    import json
    auth = pilot.auth
    roots, _ = library.workspace_folders(auth)
    reader = pilot.client.app.state.authenticator.verify('assistant')
    colleague = pilot.client.app.state.authenticator.verify('colleague')
    original = library.context(reader)['revision']
    private = library.upload(auth, 'Secret_SKILL.md', b'Secret instructions.', visibility='private')
    assert all(n['id'] != private for n in library.snapshot(colleague))
    reduced = AuthContext(auth.principal.model_copy(update={'sensitivities': ['public']}))
    assert library.context(reduced)['folders'] == []
    node = library.upload(auth, 'Humanizer_SKILL.md', b'# Humanizer\nUse clear language.', proposed=True)
    assert all(n['id'] != node for n in library.snapshot(reader))
    assert any(n['id'] == node for n in library.snapshot(auth, True))
    library.review(auth, node, True)
    assert library.info(auth, node)['parent_id'] == roots['Skills']
    approved = library.context(reader)['revision']
    assert original != approved
    assert process_one(library, auth)
    while process_one(library, auth):
        pass
    indexed = library.context(reader)['revision']
    assert approved != indexed
    assert library.context(reader)['counts']['skills'] == 1
    revision = library.context(colleague)['revision']
    library.upload(auth, 'Secret_SKILL.md', b'Changed secret.', node_id=private)
    assert library.context(colleague)['revision'] == revision
    changed = library.upload(auth, 'Humanizer_SKILL.md', b'# Humanizer\nUse shorter sentences.', node_id=node)
    assert changed == node and library.context(reader)['revision'] != indexed
    def call(name, arguments):
        response = pilot.client.post('/mcp/', headers={'Authorization': 'Bearer assistant',
            'Accept': 'application/json, text/event-stream'}, json={'jsonrpc': '2.0', 'id': 33,
            'method': 'tools/call', 'params': {'name': name, 'arguments': arguments}})
        result = response.json()['result']
        assert not result.get('isError'), result
        if name == 'list_skills':
            return [item for block in result['content']
                    for value in [json.loads(block['text'])]
                    for item in (value if isinstance(value, list) else [value])]
        return json.loads(result['content'][0]['text'])
    catalog = call('list_skills', {})
    assert any(item['name'] == 'Humanizer_SKILL.md' and item['version'] == '2' for item in catalog)
    exact = call('get_skill', {'name': 'Humanizer_SKILL.md', 'version': '1'})
    assert exact['definition'] == '# Humanizer\nUse clear language.'
    assert call('get_workspace_context', {})['counts']['skills'] == 1
    assert library.context(auth, True)['counts']['pending'] == 0


def test_spreadsheet_mcp_bounds_and_sheet_names(pilot,library):
    import io
    wb=Workbook();wb.active.title='Costs'
    for sheet in [wb.active,wb.create_sheet('Forecast')]:
        for n in range(25):sheet.append(['Heating',n])
    output=io.BytesIO();wb.save(output)
    node=library.upload(pilot.auth,'Forecast.xlsx',output.getvalue())
    assert process_one(library,pilot.auth)
    assert library.info(pilot.auth,node)['sheets']==['Costs','Forecast']
    response=pilot.client.post('/mcp/',headers={'Authorization':'Bearer assistant','Accept':'application/json, text/event-stream'},json={
        'jsonrpc':'2.0','id':22,'method':'tools/call','params':{'name':'read_spreadsheet_rows','arguments':{'file_id':str(node),'first_row':1,'last_row':100}}})
    assert response.status_code==200
    result=response.json()['result']
    assert not result.get('isError'),result
    import json
    payload=json.loads(result['content'][0]['text'])
    assert len(payload['rows'])==20
